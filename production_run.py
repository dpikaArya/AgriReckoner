"""
AAIF Production Run — End-to-End Pipeline Execution

Executes all 12 steps:
 1. External Data Source Synchronization
 2. AI Extraction Pipeline (23 agents via Orchestrator)
 3. UAMS Schema Update
 4. Observation Store Update
 5. Feature Engineering
 6. Model Drift Detection
 7. Adaptive Model Training
 8. Model Evaluation
 9. Ready Reckoner Generation
10. Framework Benchmark
11. Framework Assessment
12. Deliverable Generation

Usage:
    python production_run.py
"""

import json
import logging
import os
import shutil
import sys
import time
import traceback
from datetime import datetime
from pathlib import Path
from typing import Optional

import pandas as pd

# ── Base paths ──────────────────────────────────────────────────────────
BASE_DIR = Path(__file__).parent.resolve()
sys.path.insert(0, str(BASE_DIR))

from agri_ai_agent.config.settings import AgriAISettings
from agri_ai_agent.external_data.registry_db import DatasetRegistry
from src.utils.logging_config import setup_logging

# ── Settings ────────────────────────────────────────────────────────────
settings = AgriAISettings()
settings.INCREMENTAL_MODE = True
settings.CHECKPOINT_ENABLED = True
settings.ensure_dirs()

# Override REPORTS_DIR to project root (production_run.py output goes here,
# separate from agri_ai_agent/outputs/reports used by the package)
settings.REPORTS_DIR = BASE_DIR / "reports"
REPORTS_DIR = settings.REPORTS_DIR
REPORTS_DIR.mkdir(parents=True, exist_ok=True)

RUN_ID = datetime.now().strftime("PROD_%Y%m%d_%H%M%S")
setup_logging(log_dir=settings.LOG_DIR)
logger = logging.getLogger("ProductionRun")

# ── Global state ────────────────────────────────────────────────────────
production_state = {
    "run_id": RUN_ID,
    "started_at": datetime.now().isoformat(),
    "steps": {},
    "errors": [],
    "status": "running",
    "completed_at": None,
}


def _save_state():
    (REPORTS_DIR / "production_state.json").write_text(
        json.dumps(production_state, indent=2, default=str), encoding="utf-8"
    )


def _record_step(name: str, status: str, details: dict = None):
    production_state["steps"][name] = {
        "status": status,
        "completed_at": datetime.now().isoformat(),
        "details": details or {},
    }
    _save_state()


def _load_master_datasets() -> pd.DataFrame:
    master_dir = settings.DATA_DIR
    all_frames = []
    if master_dir.exists():
        for f in sorted(master_dir.glob("*.xlsx")):
            try:
                df = pd.read_excel(f)
                df["_source_file"] = f.name
                all_frames.append(df)
                logger.info("  Loaded %s: %d rows x %d cols", f.name, len(df), len(df.columns))
            except Exception as e:
                logger.warning("  Failed to load %s: %s", f.name, e)
    combined = pd.concat(all_frames, ignore_index=True, sort=False) if all_frames else pd.DataFrame()
    logger.info("Combined master dataset: %d rows x %d cols", len(combined), len(combined.columns))
    return combined


# ══════════════════════════════════════════════════════════════════════════
# STEP 1 — External Data Source Synchronization
# ══════════════════════════════════════════════════════════════════════════
def step1_sync_external_data_sources() -> dict:
    logger.info("=" * 60)
    logger.info("STEP 1: Synchronize External Data Sources")
    logger.info("=" * 60)

    results = {
        "sources_configured": [],
        "sources_healthy": [],
        "sources_unhealthy": [],
        "datasets_downloaded": 0,
        "datasets_valid": 0,
        "datasets_skipped": 0,
        "total_duration_sec": 0.0,
        "registry_summary": {},
        "packages_by_source": {},
        "run_logs": [],
    }

    registry_path = settings.DATASET_REGISTRY_PATH
    download_dir = settings.EXTERNAL_DATA_DIR

    registry = DatasetRegistry(registry_path)
    logger.info("Dataset registry initialized at %s", registry_path)

    try:
        from agri_ai_agent.external_data.connector_manager import ConnectorManager

        manager = ConnectorManager(
            registry=registry,
            download_dir=download_dir,
            max_workers=2,
            retry_max_attempts=1,
            retry_base_delay=1.0,
            log_dir=settings.LOG_DIR / "connector_logs",
        )

        available = manager.list_sources()
        results["sources_configured"] = available
        logger.info("Discovered %d configured sources: %s", len(available), available)

        # Health checks
        health_results = manager.health_checks()
        healthy = sorted([s for s, h in health_results.items() if h.is_healthy])
        unhealthy = sorted([s for s, h in health_results.items() if not h.is_healthy])
        results["sources_healthy"] = healthy
        results["sources_unhealthy"] = unhealthy
        logger.info("Healthy: %d | Unhealthy: %d", len(healthy), len(unhealthy))
        for s, h in health_results.items():
            logger.info("  %s: healthy=%s, latency=%.1fms, error=%s",
                       s, h.is_healthy, h.latency_ms, h.error or "none")

        # Detect updates
        updates = manager.detect_all_updates()
        if updates:
            logger.info("Updates available: %s", json.dumps(updates, indent=2))
        else:
            logger.info("No updates detected for any source")

        if not healthy:
            logger.warning("No healthy sources to sync — using existing cached datasets")
            results["note"] = "No healthy sources; all datasets skipped"
            packages_by_source, run_logs = {}, []
        else:
            max_discovery = int(os.environ.get("AGRI_MAX_DISCOVERY_PER_SOURCE", "50"))
            quick_sources = []
            for s in healthy:
                h = health_results.get(s)
                if h and h.discovered_count <= max_discovery:
                    quick_sources.append(s)
                else:
                    logger.info("Skipping %s (%d datasets — exceeds AGRI_MAX_DISCOVERY_PER_SOURCE=%d)",
                                s, h.discovered_count if h else -1, max_discovery)
            if not quick_sources:
                logger.warning("No quick-sync sources available after filtering")
                results["note"] = "No quick-sync sources; all filtered out"
                packages_by_source, run_logs = {}, []
            else:
                t0 = time.perf_counter()
                packages_by_source, run_logs = manager.run_all(sources=quick_sources)
                elapsed = time.perf_counter() - t0

            results["datasets_downloaded"] = sum(len(v) for v in packages_by_source.values())
            results["datasets_valid"] = sum(
                1 for pkgs in packages_by_source.values() for p in pkgs if p.is_valid
            )
            results["total_duration_sec"] = round(elapsed, 2)
            results["packages_by_source"] = {
                src: [p.to_dict() for p in pkgs]
                for src, pkgs in packages_by_source.items()
            }
            results["run_logs"] = [
                {
                    "source": l.source_name,
                    "duration_sec": l.duration_sec,
                    "datasets_found": l.datasets_found,
                    "datasets_downloaded": l.datasets_downloaded,
                    "datasets_valid": l.datasets_valid,
                    "success": l.success,
                    "errors": l.errors,
                }
                for l in run_logs
            ]

            summary = manager.summarize_runs(run_logs)
            logger.info("Sync summary: %s", json.dumps(summary, indent=2))

            merged_df = manager.merge_packages(packages_by_source)
            if not merged_df.empty:
                merged_path = REPORTS_DIR / "external_data_merged.csv"
                merged_df.to_csv(merged_path, index=False)
                logger.info("Merged %d rows of external data → %s", len(merged_df), merged_path)

        results["registry_summary"] = registry.summary()
        logger.info("Registry summary: %s", json.dumps(results["registry_summary"], indent=2))

    except Exception as e:
        logger.error("Step 1 failed: %s", e)
        logger.error(traceback.format_exc())
        results["error"] = str(e)

    registry.close()
    return results


# ══════════════════════════════════════════════════════════════════════════
# STEPS 2-9 — Main Pipeline via Orchestrator
# ══════════════════════════════════════════════════════════════════════════
def steps2_9_run_pipeline(external_packages: list = None) -> dict:
    logger.info("=" * 60)
    logger.info("STEPS 2-9: Execute AI Pipeline via Orchestrator")
    logger.info("=" * 60)

    results = {
        "pipeline_id": None,
        "status": None,
        "completed_agents": [],
        "failed_agents": [],
        "agent_timings": {},
        "output_rows": 0,
        "output_columns": 0,
        "duration_sec": 0.0,
        "pipeline_results": {},
    }

    try:
        from agri_ai_agent.orchestrator import Orchestrator

        orchestrator = Orchestrator(settings=settings)
        pipeline_id = orchestrator.state.pipeline_id
        results["pipeline_id"] = pipeline_id
        logger.info("Pipeline ID: %s", pipeline_id)

        # Load all master datasets as combined input
        combined_df = _load_master_datasets()

        if combined_df.empty:
            # Fallback: create a minimal DataFrame with expected columns
            logger.warning("No master datasets found — creating pipeline with minimal data")
            combined_df = pd.DataFrame({"pipeline_input": [1]})

        # Enable incremental mode so checkpoints are honored
        settings.INCREMENTAL_MODE = True

        # Pre-create the external_data checkpoint if it doesn't already exist
        # (Step 1 may have already saved it with merged external data).
        ckpt_dir = settings.CHECKPOINT_DIR
        ckpt_dir.mkdir(parents=True, exist_ok=True)
        ext_ckpt = ckpt_dir / "external_data.parquet"
        if not ext_ckpt.exists():
            safe_df = combined_df.copy()
            for col in safe_df.columns:
                if safe_df[col].dtype == "object" or str(safe_df[col].dtype) == "string":
                    safe_df[col] = safe_df[col].astype(str)
            safe_df.to_parquet(ext_ckpt, index=False)
            logger.info("Pre-created external_data checkpoint (%d rows) — pipeline will skip it", len(combined_df))
        else:
            logger.info("external_data checkpoint already exists — reusing")

        # Clear all OTHER checkpoints to force a fresh run on remaining steps
        if ckpt_dir.exists():
            for f in ckpt_dir.glob("*.parquet"):
                if f.name != "external_data.parquet":
                    f.unlink()
                    logger.info("Cleared checkpoint: %s", f.name)

        t0 = time.perf_counter()
        result_df = orchestrator.run(df=combined_df, filepath=None, papers_dir=None)
        elapsed = time.perf_counter() - t0

        results["status"] = orchestrator.state.status
        results["completed_agents"] = orchestrator.state.completed_agents
        results["failed_agents"] = orchestrator.state.failed_agents
        results["agent_timings"] = {
            k: {
                "status": v.status,
                "time_sec": v.execution_time_sec,
                "errors": v.errors,
            }
            for k, v in orchestrator.results.items()
        }
        results["output_rows"] = len(result_df) if result_df is not None else 0
        results["output_columns"] = len(result_df.columns) if result_df is not None and not result_df.empty else 0
        results["duration_sec"] = round(elapsed, 2)
        results["pipeline_results"] = {
            k: {
                "status": v.status,
                "execution_time_sec": v.execution_time_sec,
                "errors": v.errors,
                "artifacts": v.artifacts,
            }
            for k, v in orchestrator.results.items()
        }

        logger.info("Pipeline complete: status=%s, %d agents completed, %.1fs",
                    results["status"], len(results["completed_agents"]), elapsed)

        # Save pipeline output
        if result_df is not None and not result_df.empty:
            output_path = REPORTS_DIR / "pipeline_output.csv"
            result_df.to_csv(output_path, index=False)
            logger.info("Pipeline output saved: %s (%d rows)", output_path, len(result_df))

    except Exception as e:
        logger.error("Pipeline execution failed: %s", e)
        logger.error(traceback.format_exc())
        results["error"] = str(e)
        results["status"] = "failed"

    return results


# ══════════════════════════════════════════════════════════════════════════
# STEP 10 — Benchmark Framework
# ══════════════════════════════════════════════════════════════════════════
def step10_benchmark(pipeline_results: dict, step1_results: dict) -> dict:
    logger.info("=" * 60)
    logger.info("STEP 10: Benchmark Framework")
    logger.info("=" * 60)

    results = {
        "execution_time_sec": {},
        "agents_summary": {},
        "external_sync_summary": {},
        "overall_efficiency": {},
    }

    try:
        agent_timings = pipeline_results.get("agent_timings", {})
        timings_sec = {}
        for step_key, info in agent_timings.items():
            timings_sec[step_key] = info.get("time_sec", 0)
            results["execution_time_sec"][step_key] = info.get("time_sec", 0)

        completed = pipeline_results.get("completed_agents", [])
        failed = pipeline_results.get("failed_agents", [])
        total_agents = len(completed) + len(failed)

        results["agents_summary"] = {
            "total_agents": total_agents,
            "completed": len(completed),
            "failed": len(failed),
            "success_rate": round(len(completed) / max(total_agents, 1) * 100, 1),
            "total_pipeline_time_sec": pipeline_results.get("duration_sec", 0),
        }

        external = step1_results
        results["external_sync_summary"] = {
            "sources_configured": len(external.get("sources_configured", [])),
            "sources_healthy": len(external.get("sources_healthy", [])),
            "datasets_downloaded": external.get("datasets_downloaded", 0),
            "datasets_valid": external.get("datasets_valid", 0),
            "total_sync_time_sec": external.get("total_duration_sec", 0),
        }

        total_time = pipeline_results.get("duration_sec", 0) + external.get("total_duration_sec", 0)
        rows = pipeline_results.get("output_rows", 0)
        efficiency = round(rows / max(total_time, 0.001), 2) if total_time > 0 else 0

        results["overall_efficiency"] = {
            "total_execution_time_sec": round(total_time, 2),
            "pipeline_time_sec": pipeline_results.get("duration_sec", 0),
            "external_sync_time_sec": external.get("total_duration_sec", 0),
            "output_rows": rows,
            "rows_per_second": efficiency,
        }

        # Try to read existing benchmark history
        history_path = settings.OUTPUT_DIR / "benchmark_history.json"
        history = []
        if history_path.exists():
            try:
                history = json.loads(history_path.read_text(encoding="utf-8"))
            except Exception:
                logger.warning("Failed to read benchmark history, starting fresh", exc_info=True)
                history = []
        history.append({
            "run_id": RUN_ID,
            "timestamp": datetime.now().isoformat(),
            "results": results,
        })
        history_path.parent.mkdir(parents=True, exist_ok=True)
        history_path.write_text(json.dumps(history, indent=2, default=str), encoding="utf-8")
        logger.info("Benchmark history updated: %s (%d runs)", history_path, len(history))

        # Write benchmark manifest
        benchmark_manifest = {
            "run_id": RUN_ID,
            "timestamp": datetime.now().isoformat(),
            "benchmark": results,
        }
        manifest_path = REPORTS_DIR / "benchmark_manifest.json"
        manifest_path.write_text(json.dumps(benchmark_manifest, indent=2, default=str), encoding="utf-8")
        logger.info("Benchmark manifest: %s", manifest_path)

    except Exception as e:
        logger.error("Benchmark failed: %s", e)
        logger.error(traceback.format_exc())
        results["error"] = str(e)

    return results


# ══════════════════════════════════════════════════════════════════════════
# STEP 11 — Framework Assessment
# ══════════════════════════════════════════════════════════════════════════
def step11_framework_assessment(
    step1_results: dict,
    pipeline_results: dict,
    benchmark_results: dict,
) -> dict:
    logger.info("=" * 60)
    logger.info("STEP 11: Framework Assessment")
    logger.info("=" * 60)

    assessment = {
        "run_id": RUN_ID,
        "assessment_date": datetime.now().isoformat(),
        "framework_version": "2.0.0",
        "sections": {},
    }

    try:
        # Section 1: Connector Performance
        external = step1_results
        assessment["sections"]["connector_performance"] = {
            "sources_configured": len(external.get("sources_configured", [])),
            "sources_healthy": len(external.get("sources_healthy", [])),
            "sources_unhealthy": len(external.get("sources_unhealthy", [])),
            "datasets_downloaded": external.get("datasets_downloaded", 0),
            "datasets_valid": external.get("datasets_valid", 0),
            "sync_duration_sec": external.get("total_duration_sec", 0),
            "registry_summary": external.get("registry_summary", {}),
        }

        # Section 2: Pipeline Efficiency
        pipe = pipeline_results
        assessment["sections"]["pipeline_efficiency"] = {
            "pipeline_id": pipe.get("pipeline_id", ""),
            "status": pipe.get("status", "unknown"),
            "total_agents": len(pipe.get("completed_agents", [])) + len(pipe.get("failed_agents", [])),
            "agents_completed": len(pipe.get("completed_agents", [])),
            "agents_failed": len(pipe.get("failed_agents", [])),
            "pipeline_duration_sec": pipe.get("duration_sec", 0),
            "output_rows": pipe.get("output_rows", 0),
            "output_columns": pipe.get("output_columns", 0),
        }

        # Section 3: Data Quality
        assessment["sections"]["data_quality"] = {
            "master_datasets_available": len(list(settings.DATA_DIR.glob("*.xlsx"))),
            "output_rows": pipe.get("output_rows", 0),
            "output_columns": pipe.get("output_columns", 0),
            "validated_datasets_count": external.get("datasets_valid", 0),
        }

        # Section 4: Knowledge Graph Growth
        kg_files = list((settings.OUTPUT_DIR).glob("*Knowledge*")) + \
                   list((settings.OUTPUT_DIR).glob("*knowledge*"))
        assessment["sections"]["knowledge_graph"] = {
            "artifact_count": len(kg_files),
            "artifacts": [f.name for f in kg_files],
        }

        # Section 5: Observation Growth
        obs_files = list((settings.OUTPUT_DIR).glob("*Observation*")) + \
                    list((settings.OUTPUT_DIR).glob("*observation*")) + \
                    list((settings.OUTPUT_DIR).glob("*Validated_*"))
        assessment["sections"]["observations"] = {
            "observation_files": len(obs_files),
            "files": [f.name for f in obs_files],
        }

        # Section 6: Schema Coverage
        uams_files = list((settings.OUTPUT_DIR).glob("*Schema*"))
        schema_info = {}
        for f in uams_files:
            try:
                if f.suffix == ".csv":
                    df = pd.read_csv(f)
                    schema_info[f.name] = {"rows": len(df), "cols": len(df.columns)}
                elif f.suffix == ".json":
                    data = json.loads(f.read_text(encoding="utf-8"))
                    schema_info[f.name] = {"keys": list(data.keys())[:10]}
            except Exception:
                pass
        assessment["sections"]["schema_coverage"] = {
            "schema_files": len(uams_files),
            "details": schema_info,
        }

        # Section 7: Model Performance
        model_files = list((settings.OUTPUT_DIR / "models").glob("*.pkl"))
        model_metrics = {}
        metrics_file = settings.OUTPUT_DIR / "model_metrics.xlsx"
        if metrics_file.exists():
            try:
                metrics_df = pd.read_excel(metrics_file)
                model_metrics = metrics_df.to_dict(orient="records")
            except Exception:
                pass

        assessment["sections"]["model_performance"] = {
            "models_trained": len(model_files),
            "model_files": [f.name for f in model_files],
            "metrics": model_metrics,
        }

        # Section 8: Recommendation Quality
        rec_files = list((settings.OUTPUT_DIR / "recommendations").glob("*"))
        assessment["sections"]["recommendation_quality"] = {
            "recommendation_files": len(rec_files),
            "files": [f.name for f in rec_files],
            "ready_reckoner_exists": (settings.OUTPUT_DIR / "Ready_Reckoner.xlsx").exists(),
        }

        # Section 9: Overall Assessment
        benchmark = benchmark_results
        success_rate = benchmark.get("agents_summary", {}).get("success_rate", 0)
        total_time = benchmark.get("overall_efficiency", {}).get("total_execution_time_sec", 0)
        rows = pipe.get("output_rows", 0)

        scores = {
            "connector_performance": _score_connector(external),
            "pipeline_efficiency": _score_pipeline(pipe),
            "data_quality": _score_data_quality(external, pipe),
            "model_performance": _score_models(model_files),
            "overall_readiness": _score_readiness(success_rate, model_files, pipe),
        }
        assessment["sections"]["scores"] = scores
        assessment["sections"]["overall_assessment"] = {
            "agents_success_rate_pct": success_rate,
            "total_execution_time_sec": round(total_time, 2),
            "total_output_rows": rows,
            "models_available": len(model_files),
            "readiness_score": scores["overall_readiness"],
            "recommendations": _generate_recommendations(scores),
        }

        assessment_path = REPORTS_DIR / "framework_assessment.json"
        assessment_path.write_text(json.dumps(assessment, indent=2, default=str), encoding="utf-8")
        logger.info("Framework assessment written: %s", assessment_path)

    except Exception as e:
        logger.error("Framework assessment failed: %s", e)
        logger.error(traceback.format_exc())
        assessment["error"] = str(e)

    return assessment


def _score_connector(external: dict) -> float:
    total = len(external.get("sources_configured", []))
    healthy = len(external.get("sources_healthy", []))
    valid = external.get("datasets_valid", 0)
    if total == 0:
        return 0.0
    return round((healthy / total * 0.6 + min(valid / max(total, 1), 1) * 0.4) * 100, 1)


def _score_pipeline(pipe: dict) -> float:
    completed = len(pipe.get("completed_agents", []))
    failed = len(pipe.get("failed_agents", []))
    total = completed + failed
    if total == 0:
        return 0.0
    return round(completed / total * 100, 1)


def _score_data_quality(external: dict, pipe: dict) -> float:
    valid_ds = external.get("datasets_valid", 0)
    rows = pipe.get("output_rows", 0)
    score = 0.0
    if valid_ds > 0:
        score += 50.0
    if rows > 0:
        score += min(rows / 1000 * 50, 50.0)
    return round(score, 1)


def _score_models(model_files: list) -> float:
    return round(min(len(model_files) * 20, 100), 1)


def _score_readiness(success_rate: float, model_files: list, pipe: dict) -> float:
    score = success_rate * 0.4
    score += min(len(model_files) * 10, 30)
    if pipe.get("output_rows", 0) > 0:
        score += 15
    if pipe.get("status") == "completed":
        score += 15
    return round(min(score, 100), 1)


def _generate_recommendations(scores: dict) -> list:
    recs = []
    if scores.get("connector_performance", 0) < 50:
        recs.append("Improve connector health and external data source integration")
    if scores.get("pipeline_efficiency", 0) < 80:
        recs.append("Increase pipeline agent success rate; review failed agents")
    if scores.get("data_quality", 0) < 50:
        recs.append("Improve data quality validation and increase observation count")
    if scores.get("model_performance", 0) < 60:
        recs.append("Train additional prediction models for more targets")
    if scores.get("overall_readiness", 0) < 70:
        recs.append("Address pipeline failures and expand test coverage")
    if not recs:
        recs.append("Framework is production-ready. Monitor and maintain regularly.")
    return recs


# ══════════════════════════════════════════════════════════════════════════
# STEP 12 — Generate Deliverables
# ══════════════════════════════════════════════════════════════════════════
def step12_generate_deliverables(
    step1_results: dict,
    pipeline_results: dict,
    benchmark_results: dict,
    assessment: dict,
) -> dict:
    logger.info("=" * 60)
    logger.info("STEP 12: Generate Deliverables")
    logger.info("=" * 60)

    results = {
        "files_created": [],
        "errors": [],
    }

    try:
        # 12a. framework_assessment.docx
        _generate_docx(assessment, pipeline_results, benchmark_results, step1_results)
        results["files_created"].append("reports/framework_assessment.docx")

        # 12b. framework_summary.md
        _generate_summary_md(assessment, pipeline_results, benchmark_results, step1_results)
        results["files_created"].append("reports/framework_summary.md")

        # 12c. pipeline_execution_log.html
        _generate_pipeline_html(pipeline_results)
        results["files_created"].append("reports/pipeline_execution_log.html")

        # 12d. training_manifest.json
        _generate_training_manifest(pipeline_results)
        results["files_created"].append("reports/training_manifest.json")

        # 12e. Copy existing key reports
        _copy_existing_reports(results)

        logger.info("Deliverables generated: %d files", len(results["files_created"]))

    except Exception as e:
        logger.error("Deliverable generation failed: %s", e)
        logger.error(traceback.format_exc())
        results["error"] = str(e)

    return results


def _generate_docx(assessment: dict, pipeline_results: dict, benchmark_results: dict, external_results: dict):
    """Generate framework_assessment.docx"""
    docx_path = REPORTS_DIR / "framework_assessment.docx"
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT

        doc = Document()

        # Title
        title = doc.add_heading("AAIF Framework Assessment Report", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"Run ID: {RUN_ID}")
        doc.add_paragraph(f"Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph(f"Framework Version: 2.0.0")
        doc.add_paragraph("")

        # 1. Executive Summary
        doc.add_heading("1. Executive Summary", 1)
        pipe = pipeline_results
        bench = benchmark_results
        p = doc.add_paragraph()
        p.add_run(f"Pipeline Status: ").bold = True
        p.add_run(f"{pipe.get('status', 'unknown').upper()}")
        p = doc.add_paragraph()
        p.add_run(f"Pipeline Duration: ").bold = True
        p.add_run(f"{pipe.get('duration_sec', 0):.1f} seconds")
        p = doc.add_paragraph()
        p.add_run(f"Agents Completed: ").bold = True
        p.add_run(f"{len(pipe.get('completed_agents', []))}")
        p = doc.add_paragraph()
        p.add_run(f"Output Rows: ").bold = True
        p.add_run(f"{pipe.get('output_rows', 0)}")
        p = doc.add_paragraph()
        p.add_run(f"Models Available: ").bold = True
        p.add_run(f"{len(list((BASE_DIR / 'models').glob('*.pkl')))}")
        doc.add_paragraph("")

        # 2. Architecture
        doc.add_heading("2. Pipeline Architecture", 1)
        doc.add_paragraph(
            "The AAIF framework implements a 23-step agent pipeline: "
            "Repository Sync → External Data → Dataset Normalization → Dataset Ingestion Bridge → "
            "Extraction → Evidence Fusion → Ontology → Table Intelligence → Schema Population → "
            "Knowledge Integration → Knowledge → Observation Generation → Validation → "
            "Feature Store → Feature Engineering → Model Selection → Training → Prediction → "
            "Recommendation → Fuzzy → Benchmark → Explainability → Ready Reckoner"
        )
        if pipe.get("completed_agents"):
            doc.add_heading("2.1 Completed Agents", 2)
            for agent in pipe["completed_agents"]:
                doc.add_paragraph(agent, style="List Bullet")
        if pipe.get("failed_agents"):
            doc.add_heading("2.2 Failed Agents", 2)
            for agent in pipe["failed_agents"]:
                doc.add_paragraph(
                    f"{agent.get('step', 'unknown')}: {agent.get('errors', ['no info'])[0][:100]}",
                    style="List Bullet",
                )
        doc.add_paragraph("")

        # 3. Pipeline Statistics
        doc.add_heading("3. Pipeline Statistics", 1)
        stats_table = doc.add_table(rows=8, cols=2)
        stats_table.style = "Light Grid Accent 1"
        stats_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        stats_data = [
            ("Total Agents", str(len(pipe.get("completed_agents", [])) + len(pipe.get("failed_agents", [])))),
            ("Completed", str(len(pipe.get("completed_agents", [])))),
            ("Failed", str(len(pipe.get("failed_agents", [])))),
            ("Output Rows", str(pipe.get("output_rows", 0))),
            ("Output Columns", str(pipe.get("output_columns", 0))),
            ("Pipeline Duration (s)", f"{pipe.get('duration_sec', 0):.1f}"),
            ("Success Rate (%)", f"{bench.get('agents_summary', {}).get('success_rate', 0):.1f}"),
            ("Rows/sec", f"{bench.get('overall_efficiency', {}).get('rows_per_second', 0):.2f}"),
        ]
        for i, (k, v) in enumerate(stats_data):
            stats_table.rows[i].cells[0].text = k
            stats_table.rows[i].cells[1].text = str(v)
        doc.add_paragraph("")

        # 4. Model Performance
        doc.add_heading("4. Model Performance", 1)
        model_files = list((settings.OUTPUT_DIR / "models").glob("*.pkl"))
        doc.add_paragraph(f"Number of trained models: {len(model_files)}")
        for mf in model_files:
            size_kb = round(mf.stat().st_size / 1024, 1)
            doc.add_paragraph(f"  {mf.stem} ({size_kb} KB)", style="List Bullet")

        metrics_file = settings.OUTPUT_DIR / "model_metrics.xlsx"
        if metrics_file.exists():
            try:
                metrics_df = pd.read_excel(metrics_file)
                doc.add_heading("4.1 Model Metrics Table", 2)
                t = doc.add_table(rows=len(metrics_df) + 1, cols=len(metrics_df.columns))
                t.style = "Light Grid Accent 1"
                for j, col in enumerate(metrics_df.columns):
                    t.rows[0].cells[j].text = str(col)
                for i, row in metrics_df.iterrows():
                    for j, col in enumerate(metrics_df.columns):
                        t.rows[i + 1].cells[j].text = str(row.get(col, ""))
            except Exception as e:
                doc.add_paragraph(f"  (Could not load metrics: {e})")
        doc.add_paragraph("")

        # 5. External Data
        doc.add_heading("5. External Data Sources", 1)
        ext = external_results
        ext_table = doc.add_table(rows=5, cols=2)
        ext_table.style = "Light Grid Accent 1"
        ext_data = [
            ("Sources Configured", str(len(ext.get("sources_configured", [])))),
            ("Healthy Sources", str(len(ext.get("sources_healthy", [])))),
            ("Datasets Downloaded", str(ext.get("datasets_downloaded", 0))),
            ("Datasets Validated", str(ext.get("datasets_valid", 0))),
            ("Sync Duration (s)", f"{ext.get('total_duration_sec', 0):.1f}"),
        ]
        for i, (k, v) in enumerate(ext_data):
            ext_table.rows[i].cells[0].text = k
            ext_table.rows[i].cells[1].text = str(v)
        doc.add_paragraph("")

        # 6. Scores
        scores = assessment.get("sections", {}).get("scores", {})
        if scores:
            doc.add_heading("6. Framework Scores", 1)
            score_table = doc.add_table(rows=len(scores) + 1, cols=2)
            score_table.style = "Light Grid Accent 1"
            score_table.rows[0].cells[0].text = "Category"
            score_table.rows[0].cells[1].text = "Score (%)"
            for i, (k, v) in enumerate(scores.items()):
                score_table.rows[i + 1].cells[0].text = k.replace("_", " ").title()
                score_table.rows[i + 1].cells[1].text = str(v)
        doc.add_paragraph("")

        # 7. Recommendations
        recs = assessment.get("sections", {}).get("overall_assessment", {}).get("recommendations", [])
        if recs:
            doc.add_heading("7. Recommendations", 1)
            for r in recs:
                doc.add_paragraph(r, style="List Number")
        doc.add_paragraph("")

        # 8. Agent Timing Details
        doc.add_heading("8. Agent Execution Timings", 1)
        agent_timings = pipeline_results.get("agent_timings", {})
        if agent_timings:
            timing_table = doc.add_table(rows=len(agent_timings) + 1, cols=3)
            timing_table.style = "Light Grid Accent 1"
            timing_table.rows[0].cells[0].text = "Agent"
            timing_table.rows[0].cells[1].text = "Status"
            timing_table.rows[0].cells[2].text = "Time (s)"
            for i, (step_key, info) in enumerate(sorted(agent_timings.items())):
                timing_table.rows[i + 1].cells[0].text = step_key
                timing_table.rows[i + 1].cells[1].text = info.get("status", "")
                timing_table.rows[i + 1].cells[2].text = f"{info.get('time_sec', 0):.2f}"

        doc.save(docx_path)
        logger.info("DOCX report generated: %s", docx_path)
    except ImportError:
        logger.warning("python-docx not installed — skipping DOCX generation")
        # Fallback: write plain text
        (REPORTS_DIR / "framework_assessment.docx").write_text(
            "DOCX generation requires python-docx. Install with: pip install python-docx",
            encoding="utf-8",
        )
    except Exception as e:
        logger.error("DOCX generation failed: %s", e)


def _generate_summary_md(
    assessment: dict,
    pipeline_results: dict,
    benchmark_results: dict,
    step1_results: dict,
):
    pipe = pipeline_results
    bench = benchmark_results
    ext = step1_results
    scores = assessment.get("sections", {}).get("scores", {})
    recs = assessment.get("sections", {}).get("overall_assessment", {}).get("recommendations", [])

    lines = [
        f"# AAIF Framework Assessment Summary",
        f"",
        f"**Run ID:** {RUN_ID}",
        f"**Date:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"**Framework Version:** 2.0.0",
        f"",
        f"## Executive Summary",
        f"",
        f"| Metric | Value |",
        f"|--------|-------|",
        f"| Pipeline Status | {pipe.get('status', 'unknown').upper()} |",
        f"| Agents Completed | {len(pipe.get('completed_agents', []))} / {len(pipe.get('completed_agents', [])) + len(pipe.get('failed_agents', []))} |",
        f"| Agents Failed | {len(pipe.get('failed_agents', []))} |",
        f"| Pipeline Duration | {pipe.get('duration_sec', 0):.1f} s |",
        f"| Output Rows | {pipe.get('output_rows', 0)} |",
        f"| Output Columns | {pipe.get('output_columns', 0)} |",
        f"| Models Available | {len(list((BASE_DIR / 'models').glob('*.pkl')))} |",
        f"| Success Rate | {bench.get('agents_summary', {}).get('success_rate', 0):.1f}% |",
        f"",
        f"## Framework Scores",
        f"",
        f"| Category | Score (%) |",
        f"|----------|-----------|",
    ]
    for k, v in scores.items():
        lines.append(f"| {k.replace('_', ' ').title()} | {v} |")

    lines.extend([
        f"",
        f"## Recommendations",
        f"",
    ])
    if recs:
        for i, r in enumerate(recs, 1):
            lines.append(f"{i}. {r}")
    else:
        lines.append("No recommendations at this time.")

    lines.extend([
        f"",
        f"## External Data Sources",
        f"",
        f"- Sources configured: {len(ext.get('sources_configured', []))}",
        f"- Sources healthy: {len(ext.get('sources_healthy', []))}",
        f"- Datasets downloaded: {ext.get('datasets_downloaded', 0)}",
        f"- Datasets validated: {ext.get('datasets_valid', 0)}",
    ])

    lines.extend([
        f"",
        f"## Completed Agents",
        f"",
    ])
    for agent in pipe.get("completed_agents", []):
        timing = pipe.get("agent_timings", {}).get(agent, {})
        lines.append(f"- **{agent}** ({timing.get('time_sec', 0):.2f}s)")

    if pipe.get("failed_agents"):
        lines.extend([
            f"",
            f"## Failed Agents",
            f"",
        ])
        for agent in pipe["failed_agents"]:
            lines.append(f"- {agent.get('step', 'unknown')}: {agent.get('errors', [''])[0][:100]}")

    lines.extend([
        f"",
        f"## Models",
        f"",
    ])
    for mf in sorted((settings.OUTPUT_DIR / "models").glob("*.pkl")):
        lines.append(f"- {mf.stem} ({round(mf.stat().st_size / 1024, 1)} KB)")

    summary_path = REPORTS_DIR / "framework_summary.md"
    summary_path.write_text("\n".join(lines), encoding="utf-8")
    logger.info("Summary MD generated: %s", summary_path)


def _generate_pipeline_html(pipeline_results: dict):
    pipe = pipeline_results
    completed = pipe.get("completed_agents", [])
    failed = pipe.get("failed_agents", [])
    timings = pipe.get("agent_timings", {})

    agent_rows = ""
    for step_key in sorted(set(list(timings.keys()) + completed + [f.get("step", "") for f in failed])):
        info = timings.get(step_key, {})
        status = info.get("status", "unknown")
        time_sec = info.get("time_sec", 0)
        color = "#4CAF50" if status == "success" else "#f44336" if status == "failed" else "#FF9800"
        agent_rows += f"""
          <tr>
            <td>{step_key}</td>
            <td style="color:{color};font-weight:bold;">{status.upper()}</td>
            <td>{time_sec:.2f}s</td>
          </tr>"""

    failed_rows = ""
    for agent in failed:
        failed_rows += f"""
          <tr>
            <td>{agent.get('step', '')}</td>
            <td style="color:#f44336;">{agent.get('errors', [''])[0][:200]}</td>
          </tr>"""

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Pipeline Execution Log — {RUN_ID}</title>
<style>
  body {{ font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; margin: 20px; color: #333; }}
  h1 {{ color: #1a237e; }}
  table {{ border-collapse: collapse; width: 100%; margin: 20px 0; }}
  th, td {{ border: 1px solid #ddd; padding: 8px 12px; text-align: left; }}
  th {{ background-color: #1a237e; color: white; }}
  tr:nth-child(even) {{ background-color: #f5f5f5; }}
  .summary {{ background: #e8f5e9; padding: 15px; border-radius: 8px; margin: 20px 0; }}
  .failed {{ background: #ffebee; padding: 15px; border-radius: 8px; margin: 20px 0; }}
  .status {{ display: inline-block; padding: 4px 12px; border-radius: 4px; color: white; }}
  .status-pass {{ background: #4CAF50; }}
  .status-fail {{ background: #f44336; }}
</style>
</head>
<body>
<h1>AAIF Pipeline Execution Log</h1>
<p><strong>Run ID:</strong> {RUN_ID}</p>
<p><strong>Date:</strong> {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
<p><strong>Pipeline ID:</strong> {pipe.get('pipeline_id', 'N/A')}</p>
<p><strong>Status:</strong> <span class="status {'status-pass' if pipe.get('status') == 'completed' else 'status-fail'}">{pipe.get('status', 'unknown').upper()}</span></p>

<div class="summary">
  <h2>Summary</h2>
  <p>Duration: <strong>{pipe.get('duration_sec', 0):.1f}s</strong></p>
  <p>Completed Agents: <strong>{len(completed)}</strong></p>
  <p>Failed Agents: <strong>{len(failed)}</strong></p>
  <p>Output Rows: <strong>{pipe.get('output_rows', 0)}</strong></p>
  <p>Output Columns: <strong>{pipe.get('output_columns', 0)}</strong></p>
</div>

<h2>Agent Execution Log</h2>
<table>
<thead><tr><th>Agent</th><th>Status</th><th>Duration</th></tr></thead>
<tbody>{agent_rows}</tbody>
</table>

{f'<h2>Failed Agents</h2><table><thead><tr><th>Agent</th><th>Error</th></tr></thead><tbody>{failed_rows}</tbody></table>' if failed_rows else ''}

<h2>Pipeline Topology</h2>
<pre style="background:#f5f5f5;padding:15px;overflow-x:auto;">
repository_sync → external_data → dataset_normalization → dataset_ingestion_bridge → 
extraction → evidence_fusion → ontology → table_intelligence → schema_population → 
knowledge_integration → knowledge → observation_generation → validation → 
feature_store → feature → model_selection → training → prediction → 
recommendation → fuzzy → benchmark → explainability → ready_reckoner
</pre>

<p><em>Generated by AAIF Production Run</em></p>
</body>
</html>"""

    html_path = REPORTS_DIR / "pipeline_execution_log.html"
    html_path.write_text(html, encoding="utf-8")
    logger.info("Pipeline HTML log generated: %s", html_path)


def _generate_training_manifest(pipeline_results: dict):
    model_files = list((settings.OUTPUT_DIR / "models").glob("*.pkl"))
    manifest = {
        "run_id": RUN_ID,
        "timestamp": datetime.now().isoformat(),
        "pipeline_id": pipeline_results.get("pipeline_id", ""),
        "models": [],
    }
    for mf in sorted(model_files):
        manifest["models"].append({
            "name": mf.stem,
            "path": str(mf.relative_to(BASE_DIR)),
            "size_bytes": mf.stat().st_size,
            "size_kb": round(mf.stat().st_size / 1024, 1),
        })
    manifest_path = REPORTS_DIR / "training_manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2, default=str), encoding="utf-8")
    logger.info("Training manifest generated: %s", manifest_path)


def _copy_existing_reports(results: dict):
    """Copy key existing reports into reports/ for consolidated access."""
    source_files = [
        ("outputs", "Pipeline_Provenance.json"),
        ("outputs", "Pipeline_Manifest.json"),
        ("outputs", "Model_Results_All.csv"),
        ("outputs", "Ready_Reckoner.xlsx"),
        ("outputs", "Ready_Reckoner_Updated.xlsx"),
        ("outputs", "model_metrics.xlsx"),
        ("outputs", "features_dataset.csv"),
        ("outputs", "Universal_Agricultural_Schema.csv"),
        ("outputs", "Universal_Agricultural_Schema.xlsx"),
        ("outputs", "Validation_Report.xlsx"),
        ("outputs", "Feature_Engineering_Report.md"),
        ("outputs", "Schema_Mapping_Report.md"),
        ("outputs", "Benchmark_History.json"),
        (".", "benchmark_history.json"),
    ]
    for subdir, fname in source_files:
        src = BASE_DIR / subdir / fname
        if src.exists():
            dst = REPORTS_DIR / fname
            try:
                shutil.copy2(src, dst)
                results["files_created"].append(f"reports/{fname}")
                logger.debug("Copied %s → reports/", fname)
            except Exception as e:
                logger.debug("Could not copy %s: %s", fname, e)


# ══════════════════════════════════════════════════════════════════════════
# MAIN EXECUTION
# ══════════════════════════════════════════════════════════════════════════
def main():
    overall_start = time.perf_counter()
    logger.info("")
    logger.info("╔" + "═" * 58 + "╗")
    logger.info("║  AAIF PRODUCTION RUN — %s  ║", RUN_ID)
    logger.info("╚" + "═" * 58 + "╝")
    logger.info("")

    final_summary = {
        "run_id": RUN_ID,
        "started_at": production_state["started_at"],
        "steps": {},
        "overall_status": "running",
        "total_execution_time_sec": 0,
    }

    # ── STEP 1 ─────────────────────────────────────────────────────────
    try:
        logger.info("")
        logger.info("▰" * 40)
        logger.info("▰  STEP 1: External Data Source Synchronization")
        logger.info("▰" * 40)
        step1_results = step1_sync_external_data_sources()
        _record_step("step1_external_data_sync", "completed" if "error" not in step1_results else "failed", step1_results)
    except Exception as e:
        logger.error("Step 1 CRASHED: %s", e)
        step1_results = {"error": str(e), "sources_configured": [], "sources_healthy": [],
                         "sources_unhealthy": [], "datasets_downloaded": 0, "datasets_valid": 0}
        _record_step("step1_external_data_sync", "crashed", {"error": str(e)})

    final_summary["steps"]["step1_external_data_sync"] = {
        "status": production_state["steps"].get("step1_external_data_sync", {}).get("status", "unknown"),
        "sources_healthy": len(step1_results.get("sources_healthy", [])),
        "datasets_downloaded": step1_results.get("datasets_downloaded", 0),
    }

    # If Step 1 downloaded data, enrich master datasets with external data
    # via spatial/crop key-based joins and save as checkpoint.
    if step1_results.get("datasets_downloaded", 0) > 0:
        try:
            packages_by_source = step1_results.get("packages_by_source", {})
            if packages_by_source:
                from agri_ai_agent.external_data.connector_manager import ConnectorManager
                mgr = ConnectorManager(
                    registry=DatasetRegistry(settings.DATASET_REGISTRY_PATH),
                    download_dir=settings.EXTERNAL_DATA_DIR,
                )
                combined_df = _load_master_datasets()
                enriched = mgr.enrich_packages(packages_by_source, combined_df)
                if enriched is not None and not enriched.empty:
                    ckpt_dir = settings.CHECKPOINT_DIR
                    ckpt_dir.mkdir(parents=True, exist_ok=True)
                    safe_enriched = enriched.copy()
                    for col in safe_enriched.columns:
                        if safe_enriched[col].dtype == "object" or str(safe_enriched[col].dtype) == "string":
                            safe_enriched[col] = safe_enriched[col].astype(str)
                    safe_enriched.to_parquet(ckpt_dir / "external_data.parquet", index=False)
                    logger.info("Saved enriched master checkpoint with %d columns (%d original + external enrichments)",
                                len(enriched.columns), len(combined_df.columns) if not combined_df.empty else 0)
        except Exception as e:
            logger.warning("Could not create enriched checkpoint from Step 1 results: %s", e)

    # ── STEPS 2-9 ──────────────────────────────────────────────────────
    try:
        logger.info("")
        logger.info("▰" * 40)
        logger.info("▰  STEPS 2-9: AI Pipeline (23 agents)")
        logger.info("▰" * 40)
        pipeline_results = steps2_9_run_pipeline()
        _record_step("steps2_9_pipeline", "completed" if pipeline_results.get("status") == "completed" else "failed", pipeline_results)
    except Exception as e:
        logger.error("Steps 2-9 CRASHED: %s", e)
        pipeline_results = {"status": "crashed", "error": str(e), "completed_agents": [], "failed_agents": [],
                           "agent_timings": {}, "output_rows": 0, "output_columns": 0, "duration_sec": 0}
        _record_step("steps2_9_pipeline", "crashed", {"error": str(e)})

    final_summary["steps"]["steps2_9_pipeline"] = {
        "status": pipeline_results.get("status", "unknown"),
        "agents_completed": len(pipeline_results.get("completed_agents", [])),
        "agents_failed": len(pipeline_results.get("failed_agents", [])),
        "duration_sec": pipeline_results.get("duration_sec", 0),
    }

    # ── STEP 10 ────────────────────────────────────────────────────────
    try:
        logger.info("")
        logger.info("▰" * 40)
        logger.info("▰  STEP 10: Benchmark Framework")
        logger.info("▰" * 40)
        benchmark_results = step10_benchmark(pipeline_results, step1_results)
        _record_step("step10_benchmark", "completed", benchmark_results)
    except Exception as e:
        logger.error("Step 10 CRASHED: %s", e)
        benchmark_results = {"error": str(e)}
        _record_step("step10_benchmark", "crashed", {"error": str(e)})

    final_summary["steps"]["step10_benchmark"] = {
        "status": "completed" if "error" not in benchmark_results else "failed",
    }

    # ── STEP 11 ────────────────────────────────────────────────────────
    try:
        logger.info("")
        logger.info("▰" * 40)
        logger.info("▰  STEP 11: Framework Assessment")
        logger.info("▰" * 40)
        assessment = step11_framework_assessment(step1_results, pipeline_results, benchmark_results)
        _record_step("step11_assessment", "completed", {"scores": assessment.get("sections", {}).get("scores", {})})
    except Exception as e:
        logger.error("Step 11 CRASHED: %s", e)
        assessment = {"error": str(e)}
        _record_step("step11_assessment", "crashed", {"error": str(e)})

    final_summary["steps"]["step11_assessment"] = {
        "status": "completed" if "error" not in assessment else "failed",
    }

    # ── STEP 12 ────────────────────────────────────────────────────────
    try:
        logger.info("")
        logger.info("▰" * 40)
        logger.info("▰  STEP 12: Generate Deliverables")
        logger.info("▰" * 40)
        deliverables = step12_generate_deliverables(step1_results, pipeline_results, benchmark_results, assessment)
        _record_step("step12_deliverables", "completed", deliverables)
    except Exception as e:
        logger.error("Step 12 CRASHED: %s", e)
        deliverables = {"error": str(e), "files_created": []}
        _record_step("step12_deliverables", "crashed", {"error": str(e)})

    final_summary["steps"]["step12_deliverables"] = {
        "status": "completed" if "error" not in deliverables else "failed",
        "files_created": len(deliverables.get("files_created", [])),
    }

    # ── Final Summary ──────────────────────────────────────────────────
    overall_elapsed = time.perf_counter() - overall_start
    all_success = all(
        s.get("status") in ("completed", "success")
        for s in final_summary["steps"].values()
    )
    final_summary["overall_status"] = "completed" if all_success else "completed_with_errors"
    final_summary["total_execution_time_sec"] = round(overall_elapsed, 2)

    production_state["status"] = final_summary["overall_status"]
    production_state["completed_at"] = datetime.now().isoformat()
    _save_state()

    # Write final summary to reports/
    summary_path = REPORTS_DIR / "final_execution_summary.json"
    summary_path.write_text(json.dumps(final_summary, indent=2, default=str), encoding="utf-8")

    # Also write summary markdown
    _write_final_md(final_summary, step1_results, pipeline_results, assessment, deliverables)

    logger.info("")
    logger.info("╔" + "═" * 58 + "╗")
    logger.info("║  PRODUCTION RUN COMPLETE                                ║")
    logger.info("║  Status: %-40s  ║", final_summary["overall_status"])
    logger.info("║  Duration: %.1f s                                        ║", overall_elapsed)
    logger.info("║  Reports: reports/                                        ║")
    logger.info("╚" + "═" * 58 + "╝")
    logger.info("")

    print(f"\n{'='*60}")
    print(f"AAIF PRODUCTION RUN COMPLETE")
    print(f"{'='*60}")
    print(f"  Run ID:       {RUN_ID}")
    print(f"  Status:       {final_summary['overall_status']}")
    print(f"  Duration:     {overall_elapsed:.1f}s")
    print(f"  Reports:      {REPORTS_DIR}")
    print(f"{'='*60}")
    print(f"  STEP 1  External Data Sync:     {final_summary['steps'].get('step1_external_data_sync', {}).get('status', 'N/A')}")
    print(f"  STEPS 2-9 Pipeline (23 agents): {final_summary['steps'].get('steps2_9_pipeline', {}).get('status', 'N/A')}")
    print(f"  STEP 10 Benchmark:              {final_summary['steps'].get('step10_benchmark', {}).get('status', 'N/A')}")
    print(f"  STEP 11 Assessment:             {final_summary['steps'].get('step11_assessment', {}).get('status', 'N/A')}")
    print(f"  STEP 12 Deliverables:           {final_summary['steps'].get('step12_deliverables', {}).get('status', 'N/A')}")
    print(f"{'='*60}")

    return final_summary


def _write_final_md(summary, step1_results, pipeline_results, assessment, deliverables):
    pipe = pipeline_results
    ext = step1_results
    scores = assessment.get("sections", {}).get("scores", {}) if isinstance(assessment, dict) else {}

    agents_completed = len(pipe.get("completed_agents", []))
    agents_failed = len(pipe.get("failed_agents", []))
    models_count = len(list((settings.OUTPUT_DIR / "models").glob("*.pkl")))

    lines = [
        f"# AAIF Production Run — Final Execution Summary",
        f"",
        f"**Run ID:** {RUN_ID}",
        f"**Status:** {summary['overall_status']}",
        f"**Total Duration:** {summary['total_execution_time_sec']:.1f} seconds",
        f"**Date:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"",
        f"## Dataset Summary",
        f"",
        f"| Metric | Value |",
        f"|--------|-------|",
        f"| Master Datasets Available | {len(list(settings.DATA_DIR.glob('*.xlsx')))} |",
        f"| External Sources Configured | {len(ext.get('sources_configured', []))} |",
        f"| External Sources Healthy | {len(ext.get('sources_healthy', []))} |",
        f"| External Datasets Downloaded | {ext.get('datasets_downloaded', 0)} |",
        f"| External Datasets Validated | {ext.get('datasets_valid', 0)} |",
        f"",
        f"## Pipeline Statistics",
        f"",
        f"| Metric | Value |",
        f"|--------|-------|",
        f"| Pipeline ID | {pipe.get('pipeline_id', 'N/A')} |",
        f"| Pipeline Status | {pipe.get('status', 'unknown')} |",
        f"| Agents Completed | {agents_completed} / {agents_completed + agents_failed} |",
        f"| Agents Failed | {agents_failed} |",
        f"| Pipeline Duration | {pipe.get('duration_sec', 0):.1f}s |",
        f"| Output Rows | {pipe.get('output_rows', 0)} |",
        f"| Output Columns | {pipe.get('output_columns', 0)} |",
        f"",
        f"## Models",
        f"",
        f"| Model | Size |",
        f"|-------|------|",
    ]
    for mf in sorted((settings.OUTPUT_DIR / "models").glob("*.pkl")):
        lines.append(f"| {mf.stem} | {round(mf.stat().st_size / 1024, 1)} KB |")

    lines.extend([
        f"",
        f"## Framework Scores",
        f"",
        f"| Category | Score (%) |",
        f"|----------|-----------|",
    ])
    for k, v in scores.items():
        lines.append(f"| {k.replace('_', ' ').title()} | {v} |")

    lines.extend([
        f"",
        f"## Deliverables Generated",
        f"",
    ])
    for f in sorted(deliverables.get("files_created", [])):
        lines.append(f"- `{f}`")

    lines.extend([
        f"",
        f"## Completed Agents",
        f"",
    ])
    for agent in pipe.get("completed_agents", []):
        timing = pipe.get("agent_timings", {}).get(agent, {})
        lines.append(f"- **{agent}** ({timing.get('time_sec', 0):.2f}s)")

    if pipe.get("failed_agents"):
        lines.extend([
            f"",
            f"## Failed Agents",
            f"",
        ])
        for agent in pipe["failed_agents"]:
            lines.append(f"- **{agent.get('step', 'unknown')}**: {agent.get('errors', [''])[0][:200]}")

    lines.append("")
    lines.append("---")
    lines.append(f"_Generated by AAIF Production Run v2.0.0_")

    md_path = REPORTS_DIR / "final_execution_summary.md"
    md_path.write_text("\n".join(lines), encoding="utf-8")
    logger.info("Final summary written: %s", md_path)


if __name__ == "__main__":
    main()
