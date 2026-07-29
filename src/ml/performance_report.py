import json
import logging
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

logger = logging.getLogger("PerformanceReport")


class PerformanceReport:
    def __init__(self, output_dir: Optional[Path] = None):
        self.output_dir = Path(output_dir) if output_dir else Path("reports")
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.before: dict[str, Any] = {}
        self.after: dict[str, Any] = {}

    def set_baseline(self, **kwargs):
        self.before = kwargs

    def set_optimized(self, **kwargs):
        self.after = kwargs

    def generate(self) -> str:
        report = self._build_report()
        md = self._render_markdown(report)
        docx = self._render_docx(report)
        self._save(report, md, docx)
        return md

    def _build_report(self) -> dict:
        before_features = self.before.get("features", 0)
        after_features = self.after.get("features", 0)
        before_r2 = self.before.get("r2", 0)
        after_r2 = self.after.get("r2", 0)
        before_time = self.before.get("training_time_sec", 1)
        after_time = self.after.get("training_time_sec", 1)
        before_rmse = self.before.get("rmse", 1)
        after_rmse = self.after.get("rmse", 1)
        before_quality = self.before.get("data_quality", 0)
        after_quality = self.after.get("data_quality", 0)

        feature_reduction = (
            (before_features - after_features) / max(before_features, 1) * 100
            if before_features > 0
            else 0
        )
        r2_improvement = (
            (after_r2 - before_r2) / abs(before_r2) * 100
            if before_r2 != 0
            else 0
        )
        time_reduction = (
            (before_time - after_time) / max(before_time, 1) * 100
            if before_time > 0
            else 0
        )
        rmse_reduction = (
            (before_rmse - after_rmse) / max(before_rmse, 1) * 100
            if before_rmse > 0
            else 0
        )
        quality_improvement = (
            (after_quality - before_quality) / max(before_quality, 1) * 100
            if before_quality > 0
            else 0
        )

        return {
            "report_title": "AAIF Model Optimization — Efficiency Report",
            "generated_at": datetime.now().isoformat(),
            "baseline": {
                "total_features": before_features,
                "r2_score": before_r2,
                "rmse": before_rmse,
                "training_time_sec": before_time,
                "data_quality_score": before_quality,
            },
            "optimized": {
                "total_features": after_features,
                "r2_score": after_r2,
                "rmse": after_rmse,
                "training_time_sec": after_time,
                "data_quality_score": after_quality,
                "feature_importance_top_10": self.after.get("top_features", []),
                "best_model": self.after.get("best_model", ""),
                "ensemble_r2": self.after.get("ensemble_r2"),
            },
            "improvements": {
                "feature_reduction_pct": round(feature_reduction, 2),
                "r2_improvement_pct": round(r2_improvement, 2),
                "training_time_reduction_pct": round(time_reduction, 2),
                "rmse_reduction_pct": round(rmse_reduction, 2),
                "data_quality_improvement_pct": round(quality_improvement, 2),
            },
        }

    def _render_markdown(self, report: dict) -> str:
        b = report["baseline"]
        o = report["optimized"]
        imp = report["improvements"]

        return f"""# {report['report_title']}

**Generated:** {report['generated_at']}

---

## Before Optimization

| Metric | Value |
|--------|-------|
| Total Features | {b['total_features']} |
| R² Score | {b['r2_score']} |
| RMSE | {b['rmse']} |
| Training Time | {b['training_time_sec']}s |
| Data Quality Score | {b['data_quality_score']} |

## After Optimization

| Metric | Value |
|--------|-------|
| Total Features | {o['total_features']} |
| R² Score | {o['r2_score']} |
| RMSE | {o['rmse']} |
| Training Time | {o['training_time_sec']}s |
| Data Quality Score | {o['data_quality_score']} |
| Best Model | {o.get('best_model', 'N/A')} |
| Ensemble R² | {o.get('ensemble_r2', 'N/A')} |

## Improvements

| Metric | Improvement |
|--------|-------------|
| Feature Reduction | {imp['feature_reduction_pct']}% |
| R² Improvement | {imp['r2_improvement_pct']}% |
| Training Time Reduction | {imp['training_time_reduction_pct']}% |
| RMSE Reduction | {imp['rmse_reduction_pct']}% |
| Data Quality Improvement | {imp['data_quality_improvement_pct']}% |

## Top 10 Features

{chr(10).join(f'{i+1}. {f}' for i, f in enumerate(o.get('feature_importance_top_10', [])))}
"""

    def _render_docx(self, report: dict) -> Optional[str]:
        try:
            from docx import Document
            from docx.shared import Inches, Pt
        except ImportError:
            logger.info("python-docx not available; skipping .docx output")
            return None

        doc = Document()
        doc.add_heading(report["report_title"], 0)
        doc.add_paragraph(f"Generated: {report['generated_at']}")

        doc.add_heading("Before Optimization", 1)
        b = report["baseline"]
        table = doc.add_table(rows=5, cols=2)
        for i, (k, v) in enumerate(b.items()):
            table.rows[i].cells[0].text = k.replace("_", " ").title()
            table.rows[i].cells[1].text = str(v)

        doc.add_heading("After Optimization", 1)
        o = report["optimized"]
        items = [(k, v) for k, v in o.items() if not isinstance(v, list)]
        table = doc.add_table(rows=len(items), cols=2)
        for i, (k, v) in enumerate(items):
            table.rows[i].cells[0].text = k.replace("_", " ").title()
            table.rows[i].cells[1].text = str(v)

        doc.add_heading("Improvements", 1)
        imp = report["improvements"]
        table = doc.add_table(rows=len(imp), cols=2)
        for i, (k, v) in enumerate(imp.items()):
            table.rows[i].cells[0].text = k.replace("_", " ").title()
            table.rows[i].cells[1].text = f"{v}%"

        if o.get("feature_importance_top_10"):
            doc.add_heading("Top 10 Features", 1)
            for i, f in enumerate(o["feature_importance_top_10"], 1):
                doc.add_paragraph(f"{i}. {f}")

        path = self.output_dir / "AI_framework_efficiency_report.docx"
        doc.save(str(path))
        logger.info("Saved report to %s", path)
        return str(path)

    def _save(self, report: dict, md: str, docx_path: Optional[str]):
        json_path = self.output_dir / "AI_framework_efficiency_report.json"
        json_path.write_text(json.dumps(report, indent=2, default=str), encoding="utf-8")

        md_path = self.output_dir / "AI_framework_efficiency_report.md"
        md_path.write_text(md, encoding="utf-8")
        logger.info("Saved efficiency report to %s, %s", json_path, md_path)
